"""Export accepted artifacts. Markdown always; DOCX/LaTeX optional paths."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from tero.errors import TeroError


def export_markdown(source: Path, dest: Path) -> Path:
    if not source.exists():
        raise TeroError(f"No hay artefacto para exportar: {source}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
    return dest


def export_docx(source: Path, dest: Path) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise TeroError(
            "Exportar .docx requiere python-docx (pip install 'tero[docx]').",
            code="docx_optional",
        ) from exc
    markdown = source.read_text(encoding="utf-8")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    in_code = False
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            document.add_paragraph(line)
            continue
        if line.startswith("---"):
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            document.add_heading(heading.group(2), level=level)
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
            continue
        if re.match(r"^\d+\.\s+", line):
            document.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
            continue
        if line.startswith("|"):
            document.add_paragraph(line)
            continue
        if not line:
            continue
        document.add_paragraph(line)
    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(dest)
    return dest


def export_latex(
    source: Path,
    dest: Path,
    *,
    tipo: str | None = None,
    payload: dict[str, Any] | None = None,
    try_pdf: bool = False,
) -> Path:
    """Markdown/JSON → deterministic .tex template (no free-form TeX from the model)."""
    from tero.latex.render import export_latex as _export_latex

    return _export_latex(source, dest, tipo=tipo, payload=payload, try_pdf=try_pdf)


def render_latex(payload: dict[str, Any] | str, *, tipo: str | None = None) -> str:
    from tero.latex.render import render_latex as _render

    return _render(payload, tipo=tipo)
